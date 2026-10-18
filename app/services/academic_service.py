"""
Layanan Basis Data untuk Master Data (Pendataan Awal), Jadwal Pelajaran, dan Guru Piket.
Menerapkan Soft Delete, Filter Tenant Sekolah, dan Relasi Antar-Entitas.
"""

from datetime import datetime, time, timezone
from math import ceil
from typing import List, Optional, Tuple
from uuid import UUID
from sqlalchemy import func, select, update
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.models.models import (
    AcademicYear,
    Classroom,
    DayOfWeek,
    DutyTeacher,
    Schedule,
    School,
    Student,
    Subject,
    Teacher,
    User,
    UserRole,
)
from app.services.excel_service import ExcelExportService


class AcademicYearService:
    """
    Layanan pengelolaan data Tahun Pelajaran.
    """

    @staticmethod
    async def get_list(
        db: AsyncSession,
        school_id: UUID,
        page: int = 1,
        limit: int = 10,
        search: Optional[str] = None
    ) -> Tuple[List[AcademicYear], int, int]:
        query = select(AcademicYear).where(
            AcademicYear.school_id == school_id,
            AcademicYear.deleted_at.is_(None)
        )
        if search:
            query = query.where(AcademicYear.name.ilike(f"%{search.strip()}%"))

        total_res = await db.execute(select(func.count()).select_from(query.subquery()))
        total_count = total_res.scalar() or 0
        total_pages = ceil(total_count / limit) if total_count > 0 else 1

        offset = (page - 1) * limit
        items_res = await db.execute(
            query.order_by(AcademicYear.created_at.desc()).offset(offset).limit(limit)
        )
        return items_res.scalars().all(), total_count, total_pages

    @staticmethod
    async def get_active(db: AsyncSession, school_id: UUID) -> Optional[AcademicYear]:
        """
        Mengambil tahun pelajaran yang sedang aktif untuk sekolah tertentu.
        """
        query = select(AcademicYear).where(
            AcademicYear.school_id == school_id,
            AcademicYear.is_active.is_(True),
            AcademicYear.deleted_at.is_(None)
        )
        res = await db.execute(query)
        return res.scalar_one_or_none()

    @staticmethod
    async def create(
        db: AsyncSession,
        school_id: UUID,
        name: str,
        semester: str,
        is_active: bool,
        start_date=None,
        end_date=None
    ) -> AcademicYear:
        # Jika diset aktif, nonaktifkan tahun pelajaran lain di sekolah yang sama
        if is_active:
            await db.execute(
                update(AcademicYear)
                .where(AcademicYear.school_id == school_id)
                .values(is_active=False)
            )

        ay = AcademicYear(
            school_id=school_id,
            name=name.strip(),
            semester=semester.strip(),
            is_active=is_active,
            start_date=start_date,
            end_date=end_date
        )
        db.add(ay)
        await db.commit()
        await db.refresh(ay)
        return ay

    @staticmethod
    async def set_active(db: AsyncSession, school_id: UUID, ay_id: UUID) -> bool:
        """
        Mengaktifkan satu tahun pelajaran dan menonaktifkan yang lain.
        """
        await db.execute(
            update(AcademicYear)
            .where(AcademicYear.school_id == school_id)
            .values(is_active=False)
        )
        await db.execute(
            update(AcademicYear)
            .where(AcademicYear.school_id == school_id, AcademicYear.id == ay_id)
            .values(is_active=True)
        )
        await db.commit()
        return True

    @staticmethod
    async def soft_delete(db: AsyncSession, school_id: UUID, ay_id: UUID) -> bool:
        query = select(AcademicYear).where(
            AcademicYear.id == ay_id,
            AcademicYear.school_id == school_id,
            AcademicYear.deleted_at.is_(None)
        )
        res = await db.execute(query)
        ay = res.scalar_one_or_none()
        if ay:
            ay.deleted_at = datetime.now(timezone.utc)
            if ay.is_active:
                ay.is_active = False
            await db.commit()
            return True
        return False


class TeacherService:
    """
    Layanan pengelolaan data Guru.
    """

    @staticmethod
    async def get_list(
        db: AsyncSession,
        school_id: UUID,
        page: int = 1,
        limit: int = 10,
        search: Optional[str] = None
    ) -> Tuple[List[Teacher], int, int]:
        query = select(Teacher).where(
            Teacher.school_id == school_id,
            Teacher.deleted_at.is_(None)
        )
        if search:
            query = query.where(
                (Teacher.name.ilike(f"%{search.strip()}%")) |
                (Teacher.nip.ilike(f"%{search.strip()}%"))
            )

        total_res = await db.execute(select(func.count()).select_from(query.subquery()))
        total_count = total_res.scalar() or 0
        total_pages = ceil(total_count / limit) if total_count > 0 else 1

        offset = (page - 1) * limit
        items_res = await db.execute(
            query.order_by(Teacher.name.asc()).offset(offset).limit(limit)
        )
        return items_res.scalars().all(), total_count, total_pages

    @staticmethod
    async def get_all_active(db: AsyncSession, school_id: UUID) -> List[Teacher]:
        """
        Mengambil semua guru aktif untuk opsi dropdown pilihan.
        """
        query = select(Teacher).where(
            Teacher.school_id == school_id,
            Teacher.is_active.is_(True),
            Teacher.deleted_at.is_(None)
        ).order_by(Teacher.name.asc())
        res = await db.execute(query)
        return res.scalars().all()

    @staticmethod
    async def create(
        db: AsyncSession,
        school_id: UUID,
        name: str,
        nip: Optional[str] = None,
        gender: Optional[str] = None,
        phone: Optional[str] = None,
        email: Optional[str] = None,
        address: Optional[str] = None,
        employment_status: Optional[str] = "PNS"
    ) -> Teacher:
        teacher = Teacher(
            school_id=school_id,
            name=name.strip(),
            nip=nip.strip() if nip else None,
            gender=gender,
            phone=phone.strip() if phone else None,
            email=email.strip().lower() if email else None,
            address=address.strip() if address else None,
            employment_status=employment_status,
            is_active=True
        )
        db.add(teacher)
        await db.commit()
        await db.refresh(teacher)
        return teacher

    @staticmethod
    async def soft_delete(db: AsyncSession, school_id: UUID, teacher_id: UUID) -> bool:
        query = select(Teacher).where(
            Teacher.id == teacher_id,
            Teacher.school_id == school_id,
            Teacher.deleted_at.is_(None)
        )
        res = await db.execute(query)
        t = res.scalar_one_or_none()
        if t:
            t.deleted_at = datetime.now(timezone.utc)
            t.is_active = False
            await db.commit()
            return True
        return False

    @staticmethod
    async def export_to_excel(db: AsyncSession, school_id: UUID, school_name: str):
        """
        Mengekspor seluruh data guru aktif ke dalam format Excel (.xlsx) rapih.
        """
        query = select(Teacher).where(
            Teacher.school_id == school_id,
            Teacher.deleted_at.is_(None)
        ).order_by(Teacher.name.asc())
        res = await db.execute(query)
        teachers = res.scalars().all()

        headers = [
            "No",
            "Nama Lengkap & Gelar",
            "NIP / NUPTK",
            "L/P",
            "Status Pegawai",
            "No. WhatsApp/Telepon",
            "Email",
            "Alamat Tinggal"
        ]
        rows = []
        for idx, t in enumerate(teachers, start=1):
            rows.append([
                idx,
                t.name,
                t.nip or "-",
                t.gender or "-",
                t.employment_status or "-",
                t.phone or "-",
                t.email or "-",
                t.address or "-"
            ])

        return ExcelExportService.create_styled_workbook(
            title="Laporan Data Guru dan Tenaga Pendidik",
            school_name=school_name,
            headers=headers,
            rows=rows,
            sheet_name="Data Guru"
        )

    @staticmethod
    def generate_sample_excel():
        """
        Membuat template sampel data Excel untuk impor massal guru.
        """
        headers = [
            "Nama Lengkap & Gelar (Wajib)",
            "NIP / NUPTK (Opsional)",
            "Jenis Kelamin L/P (Wajib)",
            "Status Kepegawaian (PNS/PPPK/Honorer)",
            "No Telepon/WA (Opsional)",
            "Email (Opsional)",
            "Alamat (Opsional)"
        ]
        sample_rows = [
            ["Drs. H. Ahmad Dahlan, M.Pd", "197508122000031001", "L", "PNS", "081234567890", "ahmad.dahlan@smpn.sch.id", "Jl. Merdeka No. 12"],
            ["Siti Nurjanah, S.Pd", "198402152008012004", "P", "PNS", "081298765432", "siti.nurjanah@smpn.sch.id", "Jl. Kartini No. 45"],
            ["Budi Santoso, S.Kom", "", "L", "Honorer", "085612347890", "budi.santoso@gmail.com", "Jl. Sudirman No. 8"]
        ]
        return ExcelExportService.create_styled_workbook(
            title="Template Impor Massal Data Guru",
            school_name="Panduan Format",
            headers=headers,
            rows=sample_rows,
            sheet_name="Template Guru"
        )

    @staticmethod
    async def import_from_excel(db: AsyncSession, school_id: UUID, file_bytes: bytes) -> Tuple[int, List[str]]:
        """
        Mengimpor data guru massal dari file Excel.
        Mengembalikan jumlah data sukses dan daftar pesan kesalahan jika ada baris yang gagal.
        """
        from io import BytesIO
        import openpyxl

        wb = openpyxl.load_workbook(filename=BytesIO(file_bytes), data_only=True)
        ws = wb.active

        success_count = 0
        errors = []

        # Baris data dimulai dari baris ke-5 (baris 1-2 judul kop, baris 4 header kolom)
        # Atau deteksi header jika ada kata "Nama"
        header_row_idx = 4
        for r in range(1, 10):
            val = str(ws.cell(row=r, column=1).value or "").lower()
            if "nama" in val or "wajib" in val or "nip" in val:
                header_row_idx = r
                break

        start_row = header_row_idx + 1
        for row_idx in range(start_row, ws.max_row + 1):
            name_val = ws.cell(row=row_idx, column=1).value
            if not name_val or str(name_val).strip() == "":
                continue  # Baris kosong dilewati

            name = str(name_val).strip()
            nip = str(ws.cell(row=row_idx, column=2).value or "").strip() or None
            gender_val = str(ws.cell(row=row_idx, column=3).value or "").strip().upper()
            gender = "L" if gender_val.startswith("L") else ("P" if gender_val.startswith("P") else "L")
            emp_status = str(ws.cell(row=row_idx, column=4).value or "PNS").strip()
            phone = str(ws.cell(row=row_idx, column=5).value or "").strip() or None
            email = str(ws.cell(row=row_idx, column=6).value or "").strip().lower() or None
            address = str(ws.cell(row=row_idx, column=7).value or "").strip() or None

            try:
                # Cek duplikasi NIP jika ada
                if nip:
                    existing = await db.execute(
                        select(Teacher).where(
                            Teacher.school_id == school_id,
                            Teacher.nip == nip,
                            Teacher.deleted_at.is_(None)
                        )
                    )
                    if existing.scalar_one_or_none():
                        errors.append(f"Baris {row_idx}: NIP '{nip}' sudah terdaftar.")
                        continue

                t = Teacher(
                    school_id=school_id,
                    name=name,
                    nip=nip,
                    gender=gender,
                    employment_status=emp_status,
                    phone=phone,
                    email=email,
                    address=address,
                    is_active=True
                )
                db.add(t)
                success_count += 1
            except Exception as e:
                errors.append(f"Baris {row_idx}: Gagal menyimpan ({str(e)})")

        if success_count > 0:
            await db.commit()

        return success_count, errors


class SubjectService:
    """
    Layanan pengelolaan data Mata Pelajaran (Mapel).
    """

    @staticmethod
    async def get_list(
        db: AsyncSession,
        school_id: UUID,
        page: int = 1,
        limit: int = 10,
        search: Optional[str] = None
    ) -> Tuple[List[Subject], int, int]:
        query = select(Subject).where(
            Subject.school_id == school_id,
            Subject.deleted_at.is_(None)
        )
        if search:
            query = query.where(
                (Subject.name.ilike(f"%{search.strip()}%")) |
                (Subject.code.ilike(f"%{search.strip()}%"))
            )

        total_res = await db.execute(select(func.count()).select_from(query.subquery()))
        total_count = total_res.scalar() or 0
        total_pages = ceil(total_count / limit) if total_count > 0 else 1

        offset = (page - 1) * limit
        items_res = await db.execute(
            query.order_by(Subject.name.asc()).offset(offset).limit(limit)
        )
        return items_res.scalars().all(), total_count, total_pages

    @staticmethod
    async def get_all_active(db: AsyncSession, school_id: UUID) -> List[Subject]:
        query = select(Subject).where(
            Subject.school_id == school_id,
            Subject.is_active.is_(True),
            Subject.deleted_at.is_(None)
        ).order_by(Subject.name.asc())
        res = await db.execute(query)
        return res.scalars().all()

    @staticmethod
    async def create(
        db: AsyncSession,
        school_id: UUID,
        code: str,
        name: str,
        grade_level: str = "Semua",
        description: Optional[str] = None
    ) -> Subject:
        subject = Subject(
            school_id=school_id,
            code=code.strip().upper(),
            name=name.strip(),
            grade_level=grade_level,
            description=description.strip() if description else None,
            is_active=True
        )
        db.add(subject)
        await db.commit()
        await db.refresh(subject)
        return subject

    @staticmethod
    async def soft_delete(db: AsyncSession, school_id: UUID, subject_id: UUID) -> bool:
        query = select(Subject).where(
            Subject.id == subject_id,
            Subject.school_id == school_id,
            Subject.deleted_at.is_(None)
        )
        res = await db.execute(query)
        s = res.scalar_one_or_none()
        if s:
            s.deleted_at = datetime.now(timezone.utc)
            s.is_active = False
            await db.commit()
            return True
        return False


class ClassroomService:
    """
    Layanan pengelolaan data Kelas dan penetapan Wali Kelas.
    """

    @staticmethod
    async def get_list(
        db: AsyncSession,
        school_id: UUID,
        academic_year_id: Optional[UUID] = None,
        page: int = 1,
        limit: int = 10,
        search: Optional[str] = None
    ) -> Tuple[List[Classroom], int, int]:
        query = select(Classroom).options(
            selectinload(Classroom.homeroom_teacher),
            selectinload(Classroom.academic_year)
        ).where(
            Classroom.school_id == school_id,
            Classroom.deleted_at.is_(None)
        )
        if academic_year_id:
            query = query.where(Classroom.academic_year_id == academic_year_id)
        if search:
            query = query.where(Classroom.name.ilike(f"%{search.strip()}%"))

        total_res = await db.execute(select(func.count()).select_from(query.subquery()))
        total_count = total_res.scalar() or 0
        total_pages = ceil(total_count / limit) if total_count > 0 else 1

        offset = (page - 1) * limit
        items_res = await db.execute(
            query.order_by(Classroom.grade_level.asc(), Classroom.name.asc()).offset(offset).limit(limit)
        )
        return items_res.scalars().all(), total_count, total_pages

    @staticmethod
    async def get_all_active(
        db: AsyncSession,
        school_id: UUID,
        academic_year_id: Optional[UUID] = None
    ) -> List[Classroom]:
        query = select(Classroom).where(
            Classroom.school_id == school_id,
            Classroom.is_active.is_(True),
            Classroom.deleted_at.is_(None)
        )
        if academic_year_id:
            query = query.where(Classroom.academic_year_id == academic_year_id)
        res = await db.execute(query.order_by(Classroom.grade_level.asc(), Classroom.name.asc()))
        return res.scalars().all()

    @staticmethod
    async def create(
        db: AsyncSession,
        school_id: UUID,
        academic_year_id: UUID,
        name: str,
        grade_level: str,
        homeroom_teacher_id: Optional[UUID] = None,
        capacity: int = 32
    ) -> Classroom:
        classroom = Classroom(
            school_id=school_id,
            academic_year_id=academic_year_id,
            name=name.strip(),
            grade_level=grade_level.strip(),
            homeroom_teacher_id=homeroom_teacher_id,
            capacity=capacity,
            is_active=True
        )
        db.add(classroom)
        await db.commit()
        await db.refresh(classroom)
        return classroom

    @staticmethod
    async def soft_delete(db: AsyncSession, school_id: UUID, classroom_id: UUID) -> bool:
        query = select(Classroom).where(
            Classroom.id == classroom_id,
            Classroom.school_id == school_id,
            Classroom.deleted_at.is_(None)
        )
        res = await db.execute(query)
        c = res.scalar_one_or_none()
        if c:
            c.deleted_at = datetime.now(timezone.utc)
            c.is_active = False
            await db.commit()
            return True
        return False


class ScheduleService:
    """
    Layanan pengelolaan data Jadwal Pelajaran (per Tahun Pelajaran aktif).
    """

    @staticmethod
    async def get_list(
        db: AsyncSession,
        school_id: UUID,
        academic_year_id: UUID,
        classroom_id: Optional[UUID] = None,
        day_of_week: Optional[DayOfWeek] = None,
        page: int = 1,
        limit: int = 10
    ) -> Tuple[List[Schedule], int, int]:
        query = select(Schedule).options(
            selectinload(Schedule.classroom),
            selectinload(Schedule.subject),
            selectinload(Schedule.teacher)
        ).where(
            Schedule.school_id == school_id,
            Schedule.academic_year_id == academic_year_id,
            Schedule.deleted_at.is_(None)
        )
        if classroom_id:
            query = query.where(Schedule.classroom_id == classroom_id)
        if day_of_week:
            query = query.where(Schedule.day_of_week == day_of_week)

        total_res = await db.execute(select(func.count()).select_from(query.subquery()))
        total_count = total_res.scalar() or 0
        total_pages = ceil(total_count / limit) if total_count > 0 else 1

        offset = (page - 1) * limit
        items_res = await db.execute(
            query.order_by(Schedule.day_of_week.asc(), Schedule.start_time.asc())
            .offset(offset).limit(limit)
        )
        return items_res.scalars().all(), total_count, total_pages

    @staticmethod
    async def get_all_for_print(
        db: AsyncSession,
        school_id: UUID,
        academic_year_id: UUID,
        classroom_id: Optional[UUID] = None
    ) -> List[Schedule]:
        """
        Mengambil jadwal pelajaran lengkap untuk halaman cetak.
        """
        query = select(Schedule).options(
            selectinload(Schedule.classroom),
            selectinload(Schedule.subject),
            selectinload(Schedule.teacher)
        ).where(
            Schedule.school_id == school_id,
            Schedule.academic_year_id == academic_year_id,
            Schedule.deleted_at.is_(None)
        )
        if classroom_id:
            query = query.where(Schedule.classroom_id == classroom_id)

        res = await db.execute(query.order_by(Schedule.day_of_week.asc(), Schedule.start_time.asc()))
        return res.scalars().all()

    @staticmethod
    async def get_by_teacher(
        db: AsyncSession,
        school_id: UUID,
        academic_year_id: UUID,
        teacher_id: UUID
    ) -> List[Schedule]:
        """
        Mengambil jadwal mengajar lengkap untuk guru tertentu.
        """
        query = select(Schedule).options(
            selectinload(Schedule.classroom),
            selectinload(Schedule.subject),
            selectinload(Schedule.teacher)
        ).where(
            Schedule.school_id == school_id,
            Schedule.academic_year_id == academic_year_id,
            Schedule.teacher_id == teacher_id,
            Schedule.deleted_at.is_(None)
        ).order_by(Schedule.day_of_week.asc(), Schedule.start_time.asc())
        res = await db.execute(query)
        return res.scalars().all()

    @staticmethod
    async def create(
        db: AsyncSession,
        school_id: UUID,
        academic_year_id: UUID,
        classroom_id: UUID,
        subject_id: UUID,
        teacher_id: UUID,
        day_of_week: DayOfWeek,
        start_time: time,
        end_time: time,
        period_number: Optional[int] = None,
        notes: Optional[str] = None
    ) -> Schedule:
        sched = Schedule(
            school_id=school_id,
            academic_year_id=academic_year_id,
            classroom_id=classroom_id,
            subject_id=subject_id,
            teacher_id=teacher_id,
            day_of_week=day_of_week,
            start_time=start_time,
            end_time=end_time,
            period_number=period_number,
            notes=notes.strip() if notes else None
        )
        db.add(sched)
        await db.commit()
        await db.refresh(sched)
        return sched

    @staticmethod
    async def soft_delete(db: AsyncSession, school_id: UUID, schedule_id: UUID) -> bool:
        query = select(Schedule).where(
            Schedule.id == schedule_id,
            Schedule.school_id == school_id,
            Schedule.deleted_at.is_(None)
        )
        res = await db.execute(query)
        s = res.scalar_one_or_none()
        if s:
            s.deleted_at = datetime.now(timezone.utc)
            await db.commit()
            return True
        return False


class DutyTeacherService:
    """
    Layanan pengelolaan data Guru Piket Harian.
    Dapat diatur 1 guru atau lebih per hari pada Tahun Pelajaran aktif.
    """

    @staticmethod
    async def get_list(
        db: AsyncSession,
        school_id: UUID,
        academic_year_id: UUID,
        day_of_week: Optional[DayOfWeek] = None,
        page: int = 1,
        limit: int = 10
    ) -> Tuple[List[DutyTeacher], int, int]:
        query = select(DutyTeacher).options(
            selectinload(DutyTeacher.teacher)
        ).where(
            DutyTeacher.school_id == school_id,
            DutyTeacher.academic_year_id == academic_year_id,
            DutyTeacher.deleted_at.is_(None)
        )
        if day_of_week:
            query = query.where(DutyTeacher.day_of_week == day_of_week)

        total_res = await db.execute(select(func.count()).select_from(query.subquery()))
        total_count = total_res.scalar() or 0
        total_pages = ceil(total_count / limit) if total_count > 0 else 1

        offset = (page - 1) * limit
        items_res = await db.execute(
            query.order_by(DutyTeacher.day_of_week.asc()).offset(offset).limit(limit)
        )
        return items_res.scalars().all(), total_count, total_pages

    @staticmethod
    async def get_all_for_print(
        db: AsyncSession,
        school_id: UUID,
        academic_year_id: UUID
    ) -> List[DutyTeacher]:
        query = select(DutyTeacher).options(
            selectinload(DutyTeacher.teacher)
        ).where(
            DutyTeacher.school_id == school_id,
            DutyTeacher.academic_year_id == academic_year_id,
            DutyTeacher.deleted_at.is_(None)
        ).order_by(DutyTeacher.day_of_week.asc())
        res = await db.execute(query)
        return res.scalars().all()

    @staticmethod
    async def create(
        db: AsyncSession,
        school_id: UUID,
        academic_year_id: UUID,
        teacher_id: UUID,
        day_of_week: DayOfWeek,
        duty_role: str = "Guru Piket Utama",
        notes: Optional[str] = None
    ) -> DutyTeacher:
        duty = DutyTeacher(
            school_id=school_id,
            academic_year_id=academic_year_id,
            teacher_id=teacher_id,
            day_of_week=day_of_week,
            duty_role=duty_role.strip(),
            notes=notes.strip() if notes else None
        )
        db.add(duty)
        await db.commit()
        await db.refresh(duty)
        return duty

    @staticmethod
    async def soft_delete(db: AsyncSession, school_id: UUID, duty_id: UUID) -> bool:
        query = select(DutyTeacher).where(
            DutyTeacher.id == duty_id,
            DutyTeacher.school_id == school_id,
            DutyTeacher.deleted_at.is_(None)
        )
        res = await db.execute(query)
        d = res.scalar_one_or_none()
        if d:
            d.deleted_at = datetime.now(timezone.utc)
            await db.commit()
            return True
        return False


class StudentService:
    """
    Layanan pengelolaan data Siswa SMP, termasuk impor massal dan ekspor Excel rapih.
    """

    @staticmethod
    async def get_list(
        db: AsyncSession,
        school_id: UUID,
        classroom_id: Optional[UUID] = None,
        page: int = 1,
        limit: int = 10,
        search: Optional[str] = None
    ) -> Tuple[List[Student], int, int]:
        query = select(Student).options(
            selectinload(Student.classroom),
            selectinload(Student.user)
        ).where(
            Student.school_id == school_id,
            Student.deleted_at.is_(None)
        )

        if classroom_id:
            query = query.where(Student.classroom_id == classroom_id)

        if search:
            query = query.where(
                (Student.name.ilike(f"%{search.strip()}%")) |
                (Student.nisn.ilike(f"%{search.strip()}%")) |
                (Student.nis.ilike(f"%{search.strip()}%"))
            )

        total_res = await db.execute(select(func.count()).select_from(query.subquery()))
        total_count = total_res.scalar() or 0
        total_pages = ceil(total_count / limit) if total_count > 0 else 1

        offset = (page - 1) * limit
        items_res = await db.execute(
            query.order_by(Student.name.asc()).offset(offset).limit(limit)
        )
        return items_res.scalars().all(), total_count, total_pages

    @staticmethod
    async def get_by_id(db: AsyncSession, school_id: UUID, student_id: UUID) -> Optional[Student]:
        query = select(Student).options(
            selectinload(Student.classroom),
            selectinload(Student.user)
        ).where(
            Student.id == student_id,
            Student.school_id == school_id,
            Student.deleted_at.is_(None)
        )
        res = await db.execute(query)
        return res.scalar_one_or_none()

    @staticmethod
    async def get_by_user_id(db: AsyncSession, user_id: UUID) -> Optional[Student]:
        query = select(Student).options(
            selectinload(Student.classroom),
            selectinload(Student.school)
        ).where(
            Student.user_id == user_id,
            Student.deleted_at.is_(None)
        )
        res = await db.execute(query)
        return res.scalar_one_or_none()

    @staticmethod
    async def create(
        db: AsyncSession,
        school_id: UUID,
        classroom_id: Optional[UUID],
        nisn: str,
        name: str,
        gender: Optional[str] = "L",
        phone: Optional[str] = None,
        email: Optional[str] = None,
        address: Optional[str] = None,
        nis: Optional[str] = None
    ) -> Student:
        student = Student(
            school_id=school_id,
            classroom_id=classroom_id,
            nisn=nisn.strip(),
            nis=nis.strip() if nis else None,
            name=name.strip(),
            gender=gender,
            phone=phone.strip() if phone else None,
            email=email.strip().lower() if email else None,
            address=address.strip() if address else None,
            is_active=True
        )
        db.add(student)
        await db.commit()
        await db.refresh(student)
        return student

    @staticmethod
    async def soft_delete(db: AsyncSession, school_id: UUID, student_id: UUID) -> bool:
        query = select(Student).where(
            Student.id == student_id,
            Student.school_id == school_id,
            Student.deleted_at.is_(None)
        )
        res = await db.execute(query)
        s = res.scalar_one_or_none()
        if s:
            s.deleted_at = datetime.now(timezone.utc)
            s.is_active = False
            await db.commit()
            return True
        return False

    @staticmethod
    async def export_to_excel(
        db: AsyncSession,
        school_id: UUID,
        school_name: str,
        classroom_id: Optional[UUID] = None
    ):
        """
        Mengekspor seluruh data siswa aktif ke dalam format Excel (.xlsx) dengan tabel berformat resmi.
        """
        query = select(Student).options(
            selectinload(Student.classroom)
        ).where(
            Student.school_id == school_id,
            Student.deleted_at.is_(None)
        )

        if classroom_id:
            query = query.where(Student.classroom_id == classroom_id)

        query = query.order_by(Student.name.asc())
        res = await db.execute(query)
        students = res.scalars().all()

        headers = [
            "No",
            "NISN",
            "NIS",
            "Nama Lengkap Siswa",
            "L/P",
            "Kelas / Rombel",
            "No. WhatsApp / Telp",
            "Email Siswa",
            "Alamat Tinggal"
        ]
        rows = []
        for idx, s in enumerate(students, start=1):
            class_name = s.classroom.name if s.classroom else "-"
            rows.append([
                idx,
                s.nisn,
                s.nis or "-",
                s.name,
                s.gender or "-",
                class_name,
                s.phone or "-",
                s.email or "-",
                s.address or "-"
            ])

        return ExcelExportService.create_styled_workbook(
            title="Laporan Data Induk Peserta Didik (Siswa)",
            school_name=school_name,
            headers=headers,
            rows=rows,
            sheet_name="Data Siswa"
        )

    @staticmethod
    def generate_sample_excel():
        """
        Membuat template contoh data Excel untuk impor massal siswa.
        """
        headers = [
            "NISN (10 Digit Wajib)",
            "NIS (Opsional)",
            "Nama Lengkap Siswa (Wajib)",
            "L/P (Wajib)",
            "Nama Kelas (Wajib - Misal: VII A)",
            "No Telepon/WA Orang Tua (Opsional)",
            "Alamat Tinggal (Opsional)"
        ]
        sample_rows = [
            ["0081234567", "24001", "Muhammad Farhan Pratama", "L", "VII A", "081234567890", "Jl. Melati No. 10"],
            ["0081234568", "24002", "Aisyah Putri Maharani", "P", "VII A", "081298765432", "Jl. Mawar No. 15"],
            ["0081234569", "24003", "Dimas Anggara Setiawan", "L", "VII B", "085612347890", "Jl. Anggrek No. 22"]
        ]
        return ExcelExportService.create_styled_workbook(
            title="Template Impor Massal Data Siswa",
            school_name="Panduan Format",
            headers=headers,
            rows=sample_rows,
            sheet_name="Template Siswa"
        )

    @staticmethod
    async def import_from_excel(db: AsyncSession, school_id: UUID, file_bytes: bytes) -> Tuple[int, List[str]]:
        """
        Mengimpor data siswa massal dari file Excel.
        Mencocokkan nama kelas dengan data rombel aktif di sekolah.
        """
        from io import BytesIO
        import openpyxl

        wb = openpyxl.load_workbook(filename=BytesIO(file_bytes), data_only=True)
        ws = wb.active

        success_count = 0
        errors = []

        # Ambil peta kelas aktif sekolah: nama_kelas -> id
        c_res = await db.execute(
            select(Classroom).where(
                Classroom.school_id == school_id,
                Classroom.deleted_at.is_(None)
            )
        )
        classes = {c.name.strip().upper(): c.id for c in c_res.scalars().all()}

        header_row_idx = 4
        for r in range(1, 10):
            val = str(ws.cell(row=r, column=1).value or "").lower()
            if "nisn" in val or "nama" in val or "wajib" in val:
                header_row_idx = r
                break

        start_row = header_row_idx + 1
        for row_idx in range(start_row, ws.max_row + 1):
            nisn_val = ws.cell(row=row_idx, column=1).value
            name_val = ws.cell(row=row_idx, column=3).value

            if not nisn_val or str(nisn_val).strip() == "" or not name_val or str(name_val).strip() == "":
                continue

            nisn = str(nisn_val).strip()
            nis = str(ws.cell(row=row_idx, column=2).value or "").strip() or None
            name = str(name_val).strip()
            gender_val = str(ws.cell(row=row_idx, column=4).value or "").strip().upper()
            gender = "L" if gender_val.startswith("L") else ("P" if gender_val.startswith("P") else "L")
            class_name = str(ws.cell(row=row_idx, column=5).value or "").strip().upper()
            phone = str(ws.cell(row=row_idx, column=6).value or "").strip() or None
            address = str(ws.cell(row=row_idx, column=7).value or "").strip() or None

            cid = classes.get(class_name)

            try:
                # Periksa duplikasi NISN
                existing = await db.execute(
                    select(Student).where(
                        Student.nisn == nisn,
                        Student.deleted_at.is_(None)
                    )
                )
                if existing.scalar_one_or_none():
                    errors.append(f"Baris {row_idx}: NISN '{nisn}' ({name}) sudah terdaftar.")
                    continue

                student = Student(
                    school_id=school_id,
                    classroom_id=cid,
                    nisn=nisn,
                    nis=nis,
                    name=name,
                    gender=gender,
                    phone=phone,
                    address=address,
                    is_active=True
                )
                db.add(student)
                success_count += 1
            except Exception as e:
                errors.append(f"Baris {row_idx}: Gagal menyimpan {name} ({str(e)})")

        if success_count > 0:
            await db.commit()

        return success_count, errors

